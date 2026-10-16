"""
analyze_v6.py — تحليل كل تجارب v6 وإخراج الجداول للورقة
==========================================================
يقرأ: نتائج v5 (للقراءة فقط) + results_v6/* + ردود المحكّمين البشريين.
يكتب: results_v6/report_v6.md  +  results_v6/tables/*.csv

كل قسم يُحسب فقط إن كانت بياناته موجودة، فيمكن تشغيله بعد أي تجربة.
تشغيل:
    python analyze_v6.py
"""
from __future__ import annotations

import json
import math
import re
from collections import defaultdict
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats

import common
from common import DIMS, MODELS, STRONG, TOPICS, MODEL_LABEL, composite, median_scores

TAB = common.OUT_DIR / "tables"
REPORT: list[str] = []
NaN = float("nan")


# ═════════════════════════════════════════════════════════════
#  أدوات إحصائية
# ═════════════════════════════════════════════════════════════
def kalpha(matrix: np.ndarray) -> float:
    """Krippendorff α (interval). matrix: وحدات × حكّام، NaN = مفقود."""
    try:
        import krippendorff
        return float(krippendorff.alpha(reliability_data=np.asarray(matrix, float).T,
                                        level_of_measurement="interval"))
    except ImportError:
        pass
    units = [r[~np.isnan(r)] for r in np.asarray(matrix, float)]
    units = [u for u in units if len(u) >= 2]
    n = sum(len(u) for u in units)
    if n < 2:
        return NaN
    do = sum((np.subtract.outer(u, u) ** 2).sum() / (len(u) - 1) for u in units) / n
    pooled = np.concatenate(units)
    de = (np.subtract.outer(pooled, pooled) ** 2).sum() / (n * (n - 1))
    return 1 - do / de if de > 0 else 1.0


def holm(pvals: list[float]) -> list[float]:
    m = len(pvals)
    order = np.argsort(pvals)
    adj = [NaN] * m
    running = 0.0
    for rank, i in enumerate(order):
        val = min(1.0, (m - rank) * pvals[i])
        running = max(running, val)
        adj[i] = running
    return adj


def cliffs_delta(a, b) -> float:
    a, b = np.asarray(a, float), np.asarray(b, float)
    return float(np.sign(a[:, None] - b[None, :]).mean())


def boot_ci(x, n=5000, seed=0):
    x = np.asarray(x, float)
    x = x[~np.isnan(x)]
    if len(x) < 2:
        return (NaN, NaN)
    r = np.random.default_rng(seed)
    bs = r.choice(x, (n, len(x))).mean(1)
    return tuple(np.percentile(bs, [2.5, 97.5]))


def wilcoxon_p(d) -> float:
    d = np.asarray(d, float)
    d = d[~np.isnan(d)]
    if len(d) < 5 or np.allclose(d, 0):
        return NaN
    try:
        return float(stats.wilcoxon(d, zero_method="pratt").pvalue)
    except ValueError:
        return NaN


def icc_a1(x, y) -> float:
    """ICC(A,1) — اتفاق مطلق بين قياسين."""
    m = np.column_stack([x, y]).astype(float)
    m = m[~np.isnan(m).any(1)]
    n, k = m.shape
    if n < 3:
        return NaN
    gm = m.mean()
    msr = k * ((m.mean(1) - gm) ** 2).sum() / (n - 1)
    msc = n * ((m.mean(0) - gm) ** 2).sum() / (k - 1)
    sse = ((m - m.mean(1, keepdims=True) - m.mean(0, keepdims=True) + gm) ** 2).sum()
    mse = sse / ((n - 1) * (k - 1))
    return float((msr - mse) / (msr + (k - 1) * mse + k * (msc - mse) / n))


def md_table(df: pd.DataFrame, floatfmt: str = "{:.3f}") -> str:
    cols = list(df.columns)
    lines = ["| " + " | ".join(str(c) for c in cols) + " |", "|" + "---|" * len(cols)]
    for _, r in df.iterrows():
        cells = []
        for c in cols:
            v = r[c]
            if isinstance(v, float):
                cells.append("—" if math.isnan(v) else floatfmt.format(v))
            else:
                cells.append(str(v))
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def section(title: str):
    REPORT.append(f"\n## {title}\n")
    print(f"\n=== {title}")


def emit(df: pd.DataFrame, name: str, note: str = "", fmt: str = "{:.3f}"):
    TAB.mkdir(parents=True, exist_ok=True)
    df.to_csv(TAB / f"{name}.csv", index=False)
    if note:
        REPORT.append(note + "\n")
    REPORT.append(md_table(df, fmt) + "\n")
    print(df.to_string(index=False))


# ═════════════════════════════════════════════════════════════
#  تحميل البيانات
# ═════════════════════════════════════════════════════════════
def load_v5() -> pd.DataFrame:
    rows = []
    for t in TOPICS:
        for d in common.load_v5_runs(t, with_text=False):
            pj = d.get("panel_per_judge") or {}
            r = {"topic": t, "run": d["_run"], "model": d["model_key"], "cond": d["condition"],
                 "seed": d["seed"], "chars": len(d.get("final_lecture", "")),
                 "inner_comp": composite(d.get("final_inner_scores")),
                 "inner_reported": d.get("final_inner_overall", NaN),
                 "inner_scores": d.get("final_inner_scores"),
                 "v5_panel_reported": d.get("external_overall", NaN),
                 "v5_panel_comp": composite(d.get("external_scores")),
                 "v5_n_ok": d.get("panel_n_ok", NaN)}
            for role, key in [("claude", "judge_claude"), ("kimi", "judge_kimi"), ("glm", "judge_glm")]:
                v = pj.get(key) or {}
                ok = bool(v) and not v.get("error")
                r[f"v5_{role}_scores"] = v.get("scores") if ok else None
                r[f"v5_{role}"] = composite(v.get("scores")) if ok else NaN
            rows.append(r)
    return pd.DataFrame(rows)


def load_judgements() -> pd.DataFrame:
    p = common.OUT_DIR / "judgements.jsonl"
    if not p.exists():
        return pd.DataFrame()
    st = common.JsonlStore(p)
    df = pd.DataFrame(st.values())
    # لا نخلط إصدارات الحكّام: نُبقي فقط سجلات المعرّف الحالي لكل حَكَم ثابت (claude/kimi/glm/gpt4o)
    if not df.empty and "model_id" in df:
        from judges import JUDGE_SPECS
        cur = {r: mid for r, (_, mid) in JUDGE_SPECS.items()}
        stale = df.role.isin(cur) & (df.model_id != df.role.map(cur))
        if stale.any():
            print(f"تجاهل {int(stale.sum())} حكماً بإصدار حَكَم مختلف عن الحالي")
        df = df[~stale]
    return df


def jlookup(J: pd.DataFrame, role, prompt="panel", rep=1, trunc=0, ok_only=True) -> dict:
    if J.empty:
        return {}
    s = J[(J.role == role) & (J.prompt == prompt) & (J.rep == rep) & (J.trunc == trunc)]
    if ok_only:
        s = s[s.error.isna()]
    return {(r.topic, r.run): r for r in s.itertuples()}


def build_v6_panel(V: pd.DataFrame, J: pd.DataFrame) -> pd.DataFrame:
    """لوحة v6 = Claude (v5 أو المعاد) + Kimi-K3@0.2 + GLM-5.2@0.2 — كلها ببرومبت اللوحة."""
    cl = jlookup(J, "claude")
    ki = jlookup(J, "kimi")
    gl = jlookup(J, "glm")
    out = []
    for r in V.itertuples():
        key = (r.topic, r.run)
        sc = {}
        if key in cl:
            sc["claude"] = cl[key].scores
        elif r.v5_claude_scores:
            sc["claude"] = r.v5_claude_scores
        if key in ki:
            sc["kimi"] = ki[key].scores
        if key in gl:
            sc["glm"] = gl[key].scores
        per = [{"scores": v} for v in sc.values()]
        med = median_scores(per) if per else {}
        out.append({"topic": r.topic, "run": r.run, "v6_n": len(sc),
                    "v6_claude": composite(sc.get("claude")), "v6_kimi": composite(sc.get("kimi")),
                    "v6_glm": composite(sc.get("glm")), "v6_panel": composite(med) if per else NaN,
                    **{f"v6s_{j}": sc.get(j) for j in ("claude", "kimi", "glm")}})
    return V.merge(pd.DataFrame(out), on=["topic", "run"])


# ═════════════════════════════════════════════════════════════
#  A) موثوقية اللوحة
# ═════════════════════════════════════════════════════════════
def alpha_for(df: pd.DataFrame, prefix: str) -> float:
    """α مجمّع عبر الأبعاد الثمانية (كل (محاضرة، بُعد) وحدة)."""
    rows = []
    for r in df.itertuples():
        for d in DIMS:
            row = []
            for j in ("claude", "kimi", "glm"):
                s = getattr(r, f"{prefix}{j}_scores" if prefix == "v5_" else f"v6s_{j}")
                v = float(s.get(d, 0)) if s else 0
                row.append(v if v > 0 else NaN)
            rows.append(row)
    return kalpha(np.array(rows)) if rows else NaN


def sec_reliability(P: pd.DataFrame, J: pd.DataFrame):
    if P["v6_kimi"].notna().sum() < 10:
        return
    from judges import JUDGE_SPECS
    section("A. موثوقية اللوحة")
    REPORT.append(f"لوحة v5: Claude-sonnet-4-6 + Kimi-K2.6 + GLM-5.1 (Kimi/GLM بحرارة 0.7 فعلياً).  "
                  f"لوحة v6: Claude-sonnet-4-6 + {JUDGE_SPECS['kimi'][1]} + {JUDGE_SPECS['glm'][1]} (كلها 0.2).\n")
    rows = []
    for t in list(TOPICS) + ["pooled"]:
        s = P if t == "pooled" else P[P.topic == t]
        row = {"topic": t,
               "α v5 panel (paper)": alpha_for(s, "v5_"),
               "α v6 panel": alpha_for(s, "v6"),
               "α v6 strong-only": alpha_for(s[s.model.isin(STRONG)], "v6"),
               "α v6 weak-only": alpha_for(s[~s.model.isin(STRONG)], "v6")}
        rows.append(row)
    emit(pd.DataFrame(rows), "A1_alpha",
         "α المجمّع مرتفع لأن النماذج متباعدة جداً؛ α داخل الفئة هو المقياس الصادق لمقارنات داخل النموذج.")
    pm = [{"model": MODEL_LABEL[m], **{t: alpha_for(P[(P.model == m) & (P.topic == t)], "v6") for t in TOPICS}}
          for m in MODELS]
    emit(pd.DataFrame(pm), "A2_alpha_per_model")
    # أثر الحرارة
    rows = []
    for j in ("kimi", "glm"):
        s = P[[f"v5_{j}", f"v6_{j}"]].dropna()
        if len(s) > 5:
            rows.append({"judge": j, "n": len(s), "mean(v5 − v6)": (s[f"v5_{j}"] - s[f"v6_{j}"]).mean(),
                         "Spearman": stats.spearmanr(s[f"v5_{j}"], s[f"v6_{j}"])[0],
                         "MAE": (s[f"v5_{j}"] - s[f"v6_{j}"]).abs().mean()})
    if rows:
        emit(pd.DataFrame(rows), "A3_v5_vs_v6_judge",
             "نفس المحاضرات: حَكَم v5 (إصدار قديم، حرارة 0.7) مقابل خلَفه في v6 (حرارة 0.2). "
             "الفرق يجمع أثر الإصدار والحرارة معاً؛ لا يُنسب للحرارة وحدها.")
    # ثبات إعادة القياس
    rows = []
    for j in ("claude", "kimi", "glm"):
        r2 = jlookup(J, j, rep=2)
        if len(r2) < 5:
            continue
        x, y = [], []
        for (t, run), rec in r2.items():
            p = P[(P.topic == t) & (P.run == run)]
            if len(p):
                x.append(p.iloc[0][f"v6_{j}"])
                y.append(rec.composite)
        x, y = np.array(x, float), np.array(y, float)
        ok = ~(np.isnan(x) | np.isnan(y))
        rows.append({"judge": j, "n": int(ok.sum()), "ICC(A,1)": icc_a1(x[ok], y[ok]),
                     "Pearson": stats.pearsonr(x[ok], y[ok])[0], "MAE": np.abs(x[ok] - y[ok]).mean()})
    if rows:
        # اللوحة كاملة
        rep2 = {j: jlookup(J, j, rep=2) for j in ("claude", "kimi", "glm")}
        keys = set.intersection(*[set(v) for v in rep2.values()]) if all(rep2.values()) else set()
        x, y = [], []
        for k in keys:
            p = P[(P.topic == k[0]) & (P.run == k[1])]
            if len(p):
                x.append(p.iloc[0]["v6_panel"])
                y.append(composite(median_scores([{"scores": rep2[j][k].scores} for j in rep2])))
        if len(x) > 5:
            x, y = np.array(x), np.array(y)
            rows.append({"judge": "PANEL (median)", "n": len(x), "ICC(A,1)": icc_a1(x, y),
                         "Pearson": stats.pearsonr(x, y)[0], "MAE": np.abs(x - y).mean()})
        emit(pd.DataFrame(rows), "A4_test_retest", "ثبات إعادة القياس (نفس المحاضرة تُقيَّم مرّتين).")
    # ارتباطات الحكّام
    rows = []
    for a_, b_ in [("claude", "kimi"), ("claude", "glm"), ("kimi", "glm")]:
        s = P[[f"v6_{a_}", f"v6_{b_}"]].dropna()
        rows.append({"pair": f"{a_}–{b_}", "Spearman": stats.spearmanr(s.iloc[:, 0], s.iloc[:, 1])[0]})
    emit(pd.DataFrame(rows), "A5_pairwise")
    n_incomplete = int((P.v6_n < 3).sum())
    REPORT.append(f"تشغيلات بلوحة ناقصة (<3 حكّام) في v6: **{n_incomplete}** (كانت {int((P.v5_n_ok < 3).sum())} في v5).\n")


# ═════════════════════════════════════════════════════════════
#  R) هل تصمد نتائج v5 الرئيسية مع لوحة v6؟
# ═════════════════════════════════════════════════════════════
def sec_replication(P: pd.DataFrame):
    P = P[P.v6_n >= 2].copy()
    if len(P) < 60:
        return
    section("R. صمود نتائج v5 عند تبديل اللوحة (v5 → v6)")
    rows = []
    for t in TOPICS:
        s = P[P.topic == t]
        g = s.groupby("model")[["v5_panel_comp", "v6_panel"]].mean()
        c = s.groupby("cond")[["v5_panel_comp", "v6_panel"]].mean()
        rows.append({"topic": t, "n": len(s),
                     "Spearman per-lecture": stats.spearmanr(s.v5_panel_comp, s.v6_panel, nan_policy="omit")[0],
                     "Spearman model ranking": stats.spearmanr(g.v5_panel_comp, g.v6_panel)[0],
                     "Spearman condition ranking": stats.spearmanr(c.v5_panel_comp, c.v6_panel)[0],
                     "mean shift (v6 − v5)": (s.v6_panel - s.v5_panel_comp).mean()})
    emit(pd.DataFrame(rows), "R1_panel_swap_rankings")
    # فروق الشروط مقابل No-Refine لكل نموذج: هل تبقى الإشارة نفسها؟
    rows = []
    for t in TOPICS:
        for m in MODELS:
            s = P[(P.topic == t) & (P.model == m)]
            base = s[s.cond == "no_refine"]
            if base.empty:
                continue
            for cnd in sorted(set(s.cond) - {"no_refine"}):
                x = s[s.cond == cnd]
                d5 = x.v5_panel_comp.mean() - base.v5_panel_comp.mean()
                d6 = x.v6_panel.mean() - base.v6_panel.mean()
                rows.append({"topic": t, "model": MODEL_LABEL[m], "cond": cnd, "Δ v5 panel": d5,
                             "Δ v6 panel": d6, "same sign": bool(np.sign(d5) == np.sign(d6))})
    D = pd.DataFrame(rows)
    if len(D):
        emit(D, "R2_condition_effects_v5_vs_v6",
             f"اتفاق إشارة أثر الشرط (مقابل No-Refine) بين اللوحتين: {D['same sign'].mean():.0%} من {len(D)} خلية.")


# ═════════════════════════════════════════════════════════════
#  B) الحَكَم الداخلي: النموذج أم البرومبت؟
# ═════════════════════════════════════════════════════════════
def sec_inner(P: pd.DataFrame, J: pd.DataFrame):
    g = jlookup(J, "gpt4o", prompt="panel")
    if len(g) < 10:
        return
    section("B. GPT-4o: أثر النموذج مقابل أثر البرومبت")
    P = P.copy()
    P["gpt4o_panelprompt"] = [g[(r.topic, r.run)].composite if (r.topic, r.run) in g else NaN for r in P.itertuples()]
    gn = jlookup(J, "gpt4o", prompt="inner_neutral")
    P["gpt4o_innerneutral"] = [gn[(r.topic, r.run)].composite if (r.topic, r.run) in gn else NaN for r in P.itertuples()]
    rows = []
    for t in TOPICS:
        for m in MODELS:
            s = P[(P.topic == t) & (P.model == m)]
            rows.append({"topic": t, "model": MODEL_LABEL[m], "panel v6": s.v6_panel.mean(),
                         "GPT-4o inner-prompt (v5)": s.inner_comp.mean(),
                         "GPT-4o inner-neutral": s.gpt4o_innerneutral.mean(),
                         "GPT-4o panel-prompt": s.gpt4o_panelprompt.mean()})
    T = pd.DataFrame(rows)
    emit(T, "B1_inner_vs_panel_by_model")
    rows = []
    for t in TOPICS:
        s = T[T.topic == t]
        for c in ["panel v6", "GPT-4o inner-prompt (v5)", "GPT-4o inner-neutral", "GPT-4o panel-prompt"]:
            if s[c].notna().all():
                rows.append({"topic": t, "evaluator": c, "min model mean": s[c].min(), "max model mean": s[c].max(),
                             "range": s[c].max() - s[c].min()})
    emit(pd.DataFrame(rows), "B2_range_compression",
         "إن بقي مدى GPT-4o ببرومبت اللوحة ضيّقاً ⇒ الانضغاط من النموذج نفسه؛ إن اتّسع ⇒ من البرومبت.")
    return P


# ═════════════════════════════════════════════════════════════
#  C) التوافق مع البشر
# ═════════════════════════════════════════════════════════════
def load_human() -> pd.DataFrame:
    import glob
    import openpyxl
    mp = common.human_lecture_runs()
    files = sorted(glob.glob(str(common.HUMAN_DIR / "الردود_المستلمة" / "rater_*.xlsx")))
    if not mp or not files:
        return pd.DataFrame()
    rows = []
    for f in files:
        rid = re.search(r"rater_(\w+?)(__|\.xlsx)", Path(f).name).group(1)
        ws = openpyxl.load_workbook(f, data_only=True)["Scoring"]
        hdr = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            r = dict(zip(hdr, r))
            if not r.get("Lecture"):
                continue
            vals = [r.get(d) for d in DIMS]
            if any(v is None for v in vals):
                continue
            lec = str(r["Lecture"]).strip()
            t, run = mp[lec]
            rows.append({"rater": rid, "lec": lec, "topic": t, "run": run,
                         "h": float(np.mean([float(v) for v in vals]))})
    return pd.DataFrame(rows)


def sec_human(P: pd.DataFrame, J: pd.DataFrame):
    H = load_human()
    if H.empty:
        return
    section("C. التوافق مع المحكّمين البشريين")
    REPORT.append(f"عدد الأشخاص (ملفات الردود): **{H.rater.nunique()}** — عدد المحاضرات: {H.lec.nunique()}.\n")
    L = H.groupby(["lec", "topic", "run"]).h.mean().reset_index().merge(P, on=["topic", "run"], how="left")
    g = jlookup(J, "gpt4o", prompt="panel")
    L["gpt4o_panelprompt"] = [g[(r.topic, r.run)].composite if (r.topic, r.run) in g else NaN for r in L.itertuples()]

    def panel4(r):
        sc = [getattr(r, f"v6s_{j}") for j in ("claude", "kimi", "glm")]
        k = (r.topic, r.run)
        if k in g:
            sc.append(g[k].scores)
        sc = [s for s in sc if s]
        return composite(median_scores([{"scores": s} for s in sc])) if len(sc) == 4 else NaN
    L["panel4"] = [panel4(r) for r in L.itertuples()]
    evaluators = {"panel v6 (median-composite)": "v6_panel", "panel v5 (as in paper)": "v5_panel_reported",
                  "Claude alone": "v6_claude", "Kimi-K3 alone (v6)": "v6_kimi", "GLM-5.2 alone (v6)": "v6_glm",
                  "GPT-4o inner-prompt (paper's inner)": "inner_reported",
                  "GPT-4o panel-prompt": "gpt4o_panelprompt", "panel of 4 (+GPT-4o)": "panel4"}
    rows = []
    for name, col in evaluators.items():
        for t in list(TOPICS) + ["pooled"]:
            s = L if t == "pooled" else L[L.topic == t]
            s = s[[col, "h"]].dropna()
            if len(s) < 6:
                continue
            x, h = s[col].values, s.h.values
            # MAE بعد معايرة خطية (leave-one-out) — هل الفرق مجرّد إزاحة في المقياس؟
            loo = []
            for i in range(len(x)):
                m = np.ones(len(x), bool)
                m[i] = False
                b = np.polyfit(x[m], h[m], 1)
                loo.append(abs(np.polyval(b, x[i]) - h[i]))
            rows.append({"evaluator": name, "topic": t, "n": len(s), "MAE": np.abs(x - h).mean(),
                         "MAE calibrated (LOO)": np.mean(loo), "Spearman": stats.spearmanr(x, h)[0],
                         "Pearson": stats.pearsonr(x, h)[0]})
    emit(pd.DataFrame(rows), "C1_human_alignment")
    # فترات ثقة للفروق مقابل لوحة v6
    base = "v6_panel"
    rows = []
    r_ = np.random.default_rng(0)
    for name, col in evaluators.items():
        if col == base:
            continue
        s = L[[base, col, "h"]].dropna()
        if len(s) < 10:
            continue
        dm, dr = [], []
        for _ in range(4000):
            i = r_.integers(0, len(s), len(s))
            b = s.iloc[i]
            dm.append(np.abs(b[col] - b.h).mean() - np.abs(b[base] - b.h).mean())
            dr.append(stats.spearmanr(b[col], b.h)[0] - stats.spearmanr(b[base], b.h)[0])
        rows.append({"evaluator vs panel v6": name, "ΔMAE (evaluator − panel)": np.mean(dm),
                     "ΔMAE 95% CI": f"[{np.nanpercentile(dm, 2.5):.2f}, {np.nanpercentile(dm, 97.5):.2f}]",
                     "Δρ (evaluator − panel)": np.nanmean(dr),
                     "Δρ 95% CI": f"[{np.nanpercentile(dr, 2.5):.2f}, {np.nanpercentile(dr, 97.5):.2f}]"})
    emit(pd.DataFrame(rows), "C2_human_bootstrap",
         "فترة ثقة تشمل الصفر ⇒ لا فرق دالّ إحصائياً مع هذا العدد من المحاضرات.")
    # α البشر
    rows = []
    for t in list(TOPICS) + ["pooled"]:
        s = H if t == "pooled" else H[H.topic == t]
        piv = s.pivot_table(index="lec", columns="rater", values="h")
        rows.append({"topic": t, "human α (interval)": kalpha(piv.values), "lectures": len(piv)})
    emit(pd.DataFrame(rows), "C3_human_alpha")


# ═════════════════════════════════════════════════════════════
#  D) الحكم المتقاطع: انحياز الذات بعد إزالة التساهل
# ═════════════════════════════════════════════════════════════
def sec_cross(J: pd.DataFrame):
    if J.empty or "experiment" not in J:
        return
    X = J[(J.experiment == "exp4") & J.error.isna()].copy()
    if len(X) < 30:
        return
    section("D. الحكم المتقاطع: تفضيل ذاتي حقيقي أم تساهل؟")
    X["gen"] = X.model_key
    X["lec"] = X.topic + "|" + X.run
    X["y"] = X.composite.astype(float)
    judges = sorted(X.role.unique(), key=lambda j: (not j.startswith("gen:"), j))
    # مصفوفة المتوسطات
    M = X.pivot_table(index="role", columns="gen", values="y", aggfunc="mean").reindex(index=judges, columns=MODELS)
    M.columns = [MODEL_LABEL[c] for c in M.columns]
    M = M.reset_index().rename(columns={"role": "judge ↓ / generator →"})
    emit(M, "D1_cross_matrix", "متوسط الدرجة التي أعطاها كل حَكَم (صف) لمحاضرات كل مولّد (عمود). القطر = حكم الذات.", "{:.2f}")

    def fit(df):
        lecs = sorted(df.lec.unique())
        js = sorted(df.role.unique())
        li = {l: i for i, l in enumerate(lecs)}
        ji = {j: i for i, j in enumerate(js)}
        gens = [m for m in MODELS if f"gen:{m}" in js]
        n = len(df)
        Xm = np.zeros((n, len(lecs) + len(js) - 1 + len(gens)))
        for r_i, r in enumerate(df.itertuples()):
            Xm[r_i, li[r.lec]] = 1
            j = ji[r.role]
            if j > 0:
                Xm[r_i, len(lecs) + j - 1] = 1
            if r.role == f"gen:{r.gen}":
                Xm[r_i, len(lecs) + len(js) - 1 + gens.index(r.gen)] = 1
        beta, *_ = np.linalg.lstsq(Xm, df.y.values, rcond=None)
        self_b = dict(zip(gens, beta[len(lecs) + len(js) - 1:]))
        len_b = dict(zip(js[1:], beta[len(lecs):len(lecs) + len(js) - 1]))
        len_b[js[0]] = 0.0
        return self_b, len_b, js[0]

    self_b, len_b, ref = fit(X)
    # bootstrap على المحاضرات (داخل كل مولّد)
    r_ = np.random.default_rng(1)
    lecs_by_gen = X.groupby("gen").lec.unique().to_dict()
    boots = defaultdict(list)
    for _ in range(1000):
        pick = []
        for g_, ls in lecs_by_gen.items():
            pick += list(r_.choice(ls, len(ls)))
        parts = []
        for k_, l in enumerate(pick):
            p = X[X.lec == l].copy()
            p["lec"] = f"{l}#{k_}"
            parts.append(p)
        try:
            sb, _, _ = fit(pd.concat(parts))
            for g_, v in sb.items():
                boots[g_].append(v)
        except Exception:
            continue
    # المقارنة الساذجة (طريقة v5): الذات − وسيط اللوحة على نفس النص
    pan = X[X.role.isin(["claude", "kimi", "glm"])].groupby("lec").y.median()
    rows = []
    for g_ in MODELS:
        if g_ not in self_b:
            continue
        own = X[(X.role == f"gen:{g_}") & (X.gen == g_)]
        naive = (own.set_index("lec").y - pan.reindex(own.lec).values).mean() if len(own) else NaN
        lo, hi = np.percentile(boots[g_], [2.5, 97.5]) if boots[g_] else (NaN, NaN)
        rows.append({"generator": MODEL_LABEL[g_], "naive self − panel (v5 method)": naive,
                     f"leniency vs {ref}": len_b.get(f"gen:{g_}", NaN),
                     "self-bias (controlled)": self_b[g_], "95% CI": f"[{lo:.2f}, {hi:.2f}]"})
    emit(pd.DataFrame(rows), "D2_self_bias",
         "self-bias = كم يرفع الحَكَم درجة محاضرته فوق ما يعطيه لمحاضرات الآخرين بعد ضبط الجودة. "
         "إن كانت الفترة تشمل الصفر فلا دليل على تفضيل ذاتي؛ الفرق الساذج حينها هو تساهل.", "{:.2f}")
    fails = J[(J.experiment == "exp4") & J.error.notna()].groupby("role").size()
    if len(fails):
        REPORT.append("أخطاء تحليل JSON لكل حَكَم: " + ", ".join(f"{k}: {v}" for k, v in fails.items()) + "\n")


# ═════════════════════════════════════════════════════════════
#  E) التحسين من نفس المسودة
# ═════════════════════════════════════════════════════════════
def load_exp1() -> pd.DataFrame:
    rows = []
    for t in TOPICS:
        base = common.OUT_DIR / "exp1" / t
        draft_files = sorted(base.glob("drafts*.jsonl"))
        run_files = sorted(base.glob("runs*.jsonl"))
        if not draft_files or not run_files:
            continue
        D = {}
        for f in draft_files:
            D.update(common.JsonlStore(f).records)
        for dr in D.values():
            if dr.get("error") or dr["panel"].get("n_ok", 3) < 3:
                continue
            rows.append({"topic": t, "model": dr["model_key"], "seed": dr["seed"], "cond": "no_refine",
                         "panel": dr["panel"]["composite"], "panel_scores": dr["panel"]["median_scores"],
                         "draft_panel": dr["panel"]["composite"], "draft_scores": dr["panel"]["median_scores"],
                         "chars": dr["lecture_chars"], "draft_chars": dr["lecture_chars"],
                         "rollbacks": 0, "rounds": 0, "targets": [], "n_ok": dr["panel"]["n_ok"],
                         "final_is_draft": True})
        for f in run_files:
            for r in common.JsonlStore(f).values():
                if r.get("error") or r["panel"].get("n_ok", 3) < 3:   # ناقصة ⇒ تُستبعد (وتُعاد عند التشغيل التالي)
                    continue
                from routes import LOCAL_MODELS
                if r["model_key"] not in LOCAL_MODELS and not (r.get("settings") or {}).get("guard_v2"):
                    continue                # سُجّلت قبل مراقبة عميل Together ⇒ أُعيدت
                dr = D.get(f"{t}|{r['model_key']}|seed{r['seed']}")
                if dr is None:
                    continue
                targets = sorted({d for it in r["iterations"][1:] if not it.get("rolled_back")
                                  for d in it.get("targets", [])})
                rows.append({"topic": t, "model": r["model_key"], "seed": r["seed"], "cond": r["condition"],
                             "panel": r["panel_composite"], "panel_scores": r["panel"]["median_scores"],
                             "draft_panel": r["draft_panel_composite"], "draft_scores": dr["panel"]["median_scores"],
                             "chars": r["final_chars"], "draft_chars": dr["lecture_chars"],
                             "rollbacks": r["n_rollbacks"], "rounds": r["n_rounds"], "targets": targets,
                             "n_ok": r["panel"]["n_ok"], "final_is_draft": r.get("final_is_draft", False)})
    return pd.DataFrame(rows)


def sec_exp1(E: pd.DataFrame):
    if E.empty or E.cond.nunique() < 2:
        return
    section("E. التحسين من نفس المسودة (تصميم مقترن)")
    W = E.pivot_table(index=["topic", "model", "seed"], columns="cond", values="panel")
    conds = [c for c in ["self_refine", "self_refine_guard", "datr_no_guard", "datr_full", "datr_composite_guard"] if c in W]
    comparisons = [(c, "no_refine") for c in conds]
    comparisons += [(a, b) for a, b in [("datr_full", "self_refine_guard"),      # أثر التوجيه مع ثبات الحارس
                                       ("datr_full", "datr_no_guard"),          # أثر الحارس مع ثبات التوجيه
                                       ("self_refine_guard", "self_refine"),    # أثر الحارس بلا توجيه
                                       ("datr_no_guard", "self_refine"),        # أثر التوجيه بلا حارس
                                       ("datr_composite_guard", "datr_full")]   # حارس مركّب مقابل لكل بُعد
                    if a in W and b in W]
    rows, ps = [], []
    present = [t for t in TOPICS if t in E.topic.unique()]
    for a, b in comparisons:
        for t in present + (["pooled"] if len(present) > 1 else []):
            s = W if t == "pooled" else W.xs(t, level="topic", drop_level=False)
            d = (s[a] - s[b]).dropna()
            if len(d) < 5:
                continue
            p = wilcoxon_p(d)
            lo, hi = boot_ci(d)
            rows.append({"comparison": f"{a} − {b}", "scope": t, "n pairs": len(d), "mean Δ": d.mean(),
                         "95% CI": f"[{lo:.2f}, {hi:.2f}]", "Wilcoxon p": p,
                         "Cliff δ": cliffs_delta(s[a].dropna(), s[b].dropna())})
            ps.append(p)
    T = pd.DataFrame(rows)
    if T.empty:
        REPORT.append("بيانات exp1 غير كافية بعد (أقل من 5 أزواج لكل مقارنة).\n")
        return
    T["Holm p"] = holm([p if p == p else 1.0 for p in T["Wilcoxon p"]])
    emit(T, "E1_paired_tests",
         "كل الفروق مقترنة على نفس المسودة. تصحيح Holm على كل المقارنات المعروضة معاً.")
    # لكل نموذج
    rows = []
    for t in present:
        for m in MODELS:
            try:
                s = W.loc[(t, m)]
            except KeyError:
                continue
            if "no_refine" not in s:
                continue
            row = {"topic": t, "model": MODEL_LABEL[m], "no_refine": s["no_refine"].mean()}
            for c in conds:
                row[c] = s[c].mean() if c in s else NaN
            rows.append(row)
    emit(pd.DataFrame(rows), "E2_per_model_means", fmt="{:.2f}")
    # الأبعاد المستهدفة مقابل غير المستهدفة (بحكم اللوحة)
    rows = []
    for c in conds:
        s = E[E.cond == c]
        tg, nt = [], []
        for r in s.itertuples():
            for d in DIMS:
                delta = r.panel_scores.get(d, NaN) - r.draft_scores.get(d, NaN)
                (tg if d in r.targets else nt).append(delta)
        rows.append({"condition": c, "Δ targeted dims (panel)": np.nanmean(tg) if tg else NaN,
                     "Δ non-targeted dims (panel)": np.nanmean(nt) if nt else NaN,
                     "rollback rate": s.rollbacks.sum() / max(1, s.rounds.sum()),
                     "Δ chars": (s.chars - s.draft_chars).mean(),
                     "final == draft (%)": 100 * s.final_is_draft.mean()})
    emit(pd.DataFrame(rows), "E3_targeting_mechanism",
         "هل يرفع DATR البُعد المستهدف فعلاً بحكم اللوحة؟ وما الضرر الجانبي على البقية؟", "{:.3f}")
    # لكل بُعد: DATR-full مقابل المسودة
    rows = []
    for c in conds:
        s = E[E.cond == c]
        row = {"condition": c}
        for d in DIMS:
            row[d[:12]] = np.nanmean([r.panel_scores.get(d, NaN) - r.draft_scores.get(d, NaN) for r in s.itertuples()])
        rows.append(row)
    emit(pd.DataFrame(rows), "E4_per_dimension_delta", fmt="{:+.2f}")


# ═════════════════════════════════════════════════════════════
#  F) الطول  ،  G) دقّة ALLaM  ،  H) التكلفة
# ═════════════════════════════════════════════════════════════
def sec_length(P: pd.DataFrame, J: pd.DataFrame):
    if J.empty:
        return
    X = J[(J.experiment == "exp5") & J.error.isna()]
    if len(X) < 10:
        return
    section("F. اختبار أثر الطول")
    trunc = int(X.trunc.iloc[0])
    per = defaultdict(list)
    for r in X.itertuples():
        per[(r.topic, r.run)].append({"scores": r.scores})
    rows = []
    for (t, run), sc in per.items():
        p = P[(P.topic == t) & (P.run == run)].iloc[0]
        rows.append({"topic": t, "model": p.model, "full": p.v6_panel, "truncated": composite(median_scores(sc)),
                     "full_chars": p.chars})
    T = pd.DataFrame(rows)
    out = []
    for t in TOPICS:
        for m in sorted(STRONG):
            s = T[(T.topic == t) & (T.model == m)]
            out.append({"topic": t, "model": MODEL_LABEL[m], "full length": s.full.mean(),
                        f"truncated to {trunc}": s.truncated.mean(), "drop": (s.full - s.truncated).mean()})
        for m in ("allam_7b", "llama_3_8b", "llama_3_3_70b"):
            s = P[(P.topic == t) & (P.model == m) & (P.cond == "no_refine")]
            out.append({"topic": t, "model": MODEL_LABEL[m] + " (reference, full)", "full length": s.v6_panel.mean(),
                        f"truncated to {trunc}": NaN, "drop": NaN, })
    emit(pd.DataFrame(out), "F1_length_control",
         f"إن بقيت درجات النماذج القوية بعد القصّ إلى {trunc} حرف أعلى بوضوح من ALLaM/Llama ⇒ الترتيب ليس أثراً للطول.", "{:.2f}")


def sec_allam(P: pd.DataFrame):
    p = common.OUT_DIR / "exp6_allam_precision.jsonl"
    if not p.exists():
        return
    section("G. ALLaM: دقّة كاملة مقابل 4-bit")
    X = pd.DataFrame(common.JsonlStore(p).values())
    rows = []
    for t in TOPICS:
        q4 = P[(P.topic == t) & (P.model == "allam_7b") & (P.cond == "no_refine")]
        fp = X[X.topic == t]
        rows.append({"topic": t, "Q4 (v5) panel": q4.v6_panel.mean(), "full-precision panel": fp.panel_composite.mean(),
                     "Q4 chars": q4.chars.mean(), "full-precision chars": fp.lecture_chars.mean(), "n fp": len(fp)})
    emit(pd.DataFrame(rows), "G1_allam_precision", fmt="{:.2f}")


def sec_cost(J: pd.DataFrame):
    if J.empty:
        return
    section("H. الرموز المستهلكة في الحكم (لتقدير التكلفة)")
    T = J.groupby("model_id")[["prompt_tokens", "completion_tokens"]].sum().reset_index()
    T["calls"] = J.groupby("model_id").size().values
    emit(T, "H1_tokens", fmt="{:,.0f}")


# ═════════════════════════════════════════════════════════════
def main():
    common.setup("backprop")
    REPORT.append("# تقرير تجارب التقوية v6\n")
    REPORT.append("كل الأرقام محسوبة آلياً من ملفات النتائج. لا تُنقل للورقة قبل مراجعتها.\n")
    V = load_v5()
    J = load_judgements()
    if not J.empty:
        J["trunc"] = J["trunc"].fillna(0).astype(int)
        J["rep"] = J["rep"].fillna(1).astype(int)
        if "error" not in J:
            J["error"] = None
        J["error"] = J["error"].where(J["error"].notna(), None)
    P = build_v6_panel(V, J)
    for fn, args in [(sec_reliability, (P, J)), (sec_replication, (P,)), (sec_inner, (P, J)), (sec_human, (P, J)),
                     (sec_cross, (J,)), (sec_exp1, (load_exp1(),)), (sec_length, (P, J)),
                     (sec_allam, (P,)), (sec_cost, (J,))]:
        try:
            fn(*args)
        except Exception as e:      # قسم واحد لا يوقف التقرير كله
            import traceback
            traceback.print_exc()
            REPORT.append(f"\n> ⚠️ تعذّر حساب {fn.__name__}: {type(e).__name__}: {e}\n")
    P.drop(columns=[c for c in P.columns if c.endswith("_scores") or c.startswith("v6s_")]).to_csv(
        common.OUT_DIR / "v6_panel_scores_240.csv", index=False)
    (common.OUT_DIR / "report_v6.md").write_text("\n".join(REPORT), encoding="utf-8")
    print(f"\n✓ التقرير: {common.OUT_DIR / 'report_v6.md'}\n✓ الجداول: {TAB}")


if __name__ == "__main__":
    main()
