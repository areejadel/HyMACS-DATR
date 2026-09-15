"""
build_human_eval.py — human-evaluation package
=====================================================
Builds the full package sent to the human raters (faculty):
  1) Selects one representative (median) lecture per (model x method) for both subjects.
     Methods: no_refine (baseline) + datr_full (treatment).
     = 6 models x 2 methods x 2 subjects = 24 lectures.
  2) Anonymizes them (L01..L24, shuffled) and saves a secret mapping for the researcher only.
  3) Exports each lecture as a clean, readable PDF.
  4) Builds an Excel rating form (8-dimension rubric + grounded anchors + cells).

Run:
    python experiments/build_human_eval.py
    python experiments/build_human_eval.py --conditions no_refine self_refine datr_no_guard datr_full   # 48 lectures
"""
from __future__ import annotations
import sys, json, glob, argparse, random, html, re
from pathlib import Path
from statistics import median

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from config import RUBRIC_DIMENSIONS

MODELS = ["allam_7b", "llama_3_8b", "llama_3_3_70b", "qwen3_235b", "deepseek_v3_1", "gemma_4_31b"]
TOPICS = ["backprop", "normalization"]
RUBRIC_DESC = {
    "academic_accuracy": "Are definitions, equations, and notation correct?",
    "pedagogical_effectiveness": "Does it teach effectively (scaffolding, examples, prerequisites)?",
    "clarity_organization": "Is the writing clear and the structure logical?",
    "practical_relevance": "Are concepts connected to real practice?",
    "completeness": "Are all important sub-topics covered (variants, limitations)?",
    "engagement_potential": "Does it motivate the reader (hooks, analogies)?",
    "technical_quality": "Is formatting precise (notation, code, tables)?",
    "innovation": "Are there fresh angles or cross-domain connections?",
}


def load_cell(topic, model, condition):
    """All valid runs for (topic, model, condition): [(seed, ext, lecture)]."""
    rows = []
    base = Path(__file__).resolve().parents[1] / "results" / topic / "runs"
    for f in base.glob(f"{model}__{condition}__seed*.json"):
        d = json.load(open(f, encoding="utf-8"))
        if d.get("error") or not d.get("final_lecture"):
            continue
        ext = float(d.get("external_overall") or 0)
        if ext <= 0:
            continue
        rows.append((int(d["seed"]), ext, d["final_lecture"]))
    return rows


def pick_median(rows):
    """Pick the median-scoring run (representative, not best, not worst)."""
    rows = sorted(rows, key=lambda r: r[1])
    return rows[len(rows) // 2]   # median (index 2 of 5)


# ─── simple md -> html (headings/lists/bold/code; keeps LaTeX as-is) ───
def md_to_html(md: str) -> str:
    out = []
    in_code = False
    for line in md.split("\n"):
        if line.strip().startswith("```"):
            out.append("</pre>" if in_code else "<pre>"); in_code = not in_code; continue
        if in_code:
            out.append(html.escape(line)); continue
        l = line.rstrip()
        m = re.match(r"^(#{1,4})\s+(.*)", l)
        if m:
            lvl = len(m.group(1)); out.append(f"<h{lvl+1}>{html.escape(m.group(2))}</h{lvl+1}>"); continue
        if re.match(r"^\s*[-*]\s+", l):
            out.append("<li>" + _inline(re.sub(r"^\s*[-*]\s+", "", l)) + "</li>"); continue
        if re.match(r"^\s*\d+\.\s+", l):
            out.append("<li>" + _inline(re.sub(r"^\s*\d+\.\s+", "", l)) + "</li>"); continue
        if not l.strip():
            out.append("<br>"); continue
        out.append("<p>" + _inline(l) + "</p>")
    return "\n".join(out)


def _inline(s: str) -> str:
    s = html.escape(s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*", r"<i>\1</i>", s)
    s = re.sub(r"`(.+?)`", r"<code>\1</code>", s)
    return s


LECTURE_CSS = """
@page{size:A4;margin:2cm 2.2cm;@bottom-center{content:"Lecture "counter(page)" / "counter(pages);font-size:9pt;color:#999;}}
body{font-family:'DejaVu Serif',Georgia,serif;font-size:11.5pt;line-height:1.65;color:#1a1a1a;}
h2{font-size:17pt;color:#11324d;border-bottom:2px solid #11324d;padding-bottom:4px;margin-top:4px;}
h3{font-size:13.5pt;color:#1d4e74;margin-top:18px;}
h4{font-size:12pt;color:#36454f;}
p{margin:7px 0;text-align:justify;} li{margin:4px 0;} code{font-family:'DejaVu Sans Mono',monospace;font-size:10pt;background:#f3f5f7;padding:1px 3px;}
pre{background:#f3f5f7;border:1px solid #dde3e8;border-radius:5px;padding:8px;font-size:9.5pt;overflow-wrap:break-word;white-space:pre-wrap;}
.hdr{background:#11324d;color:#fff;padding:14px 18px;border-radius:8px;margin-bottom:18px;}
.hdr .id{font-size:22pt;font-weight:bold;} .hdr .meta{font-size:10pt;opacity:.85;margin-top:3px;}
"""


def render_pdf(lecture_id, lecture_md, out_path):
    import weasyprint
    body = md_to_html(lecture_md)
    head = (f"<div class='hdr'><div class='id'>Lecture {lecture_id}</div>"
            f"<div class='meta'>Blind review · Score on the 8-dimension rubric (1-10)</div></div>")
    htmls = f"<html><head><meta charset='utf-8'><style>{LECTURE_CSS}</style></head><body>{head}{body}</body></html>"
    weasyprint.HTML(string=htmls).write_pdf(str(out_path))


def build_rating_form(ids, out_path, anchors_note=True):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = openpyxl.Workbook()
    # ── Instructions sheet ──
    ws = wb.active; ws.title = "Instructions"
    ws["A1"] = "Human Evaluation — Lecture Quality"; ws["A1"].font = Font(size=16, bold=True, color="11324D")
    rows = [
        "", "You will read anonymized university-level lectures and score each on 8 dimensions (1-10).",
        "1 = poor · 5-7 = typical · 9-10 = outstanding. Be calibrated and honest.",
        "Please read the two ANCHOR examples first (provided separately) to align your scale.",
        "Lectures are blind (no model/method shown) and in random order. Score independently.",
        "", "The 8 dimensions:",
    ]
    for i, t in enumerate(rows, start=3):
        ws[f"A{i}"] = t
    r = 3 + len(rows)
    for dim in RUBRIC_DIMENSIONS:
        ws[f"A{r}"] = f"• {dim}"; ws[f"A{r}"].font = Font(bold=True)
        ws[f"B{r}"] = RUBRIC_DESC[dim]; r += 1
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 70
    if anchors_note:
        ws[f"A{r+1}"] = "Anchors: see the grounded-examples folder (one excellent + one weak, pre-scored)."

    # ── Scoring sheet ──
    sc = wb.create_sheet("Scoring")
    hdr = ["Rater name", "Lecture"] + RUBRIC_DIMENSIONS + ["overall (auto)", "Comments"]
    fill = PatternFill("solid", fgColor="11324D"); bold = Font(bold=True, color="FFFFFF")
    thin = Border(*[Side(style="thin", color="CCCCCC")] * 4)
    for c, h in enumerate(hdr, start=1):
        cell = sc.cell(row=1, column=c, value=h); cell.fill = fill; cell.font = bold
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, lid in enumerate(ids, start=2):
        sc.cell(row=i, column=2, value=lid).font = Font(bold=True)
        # overall = average of the 8 dimension columns (C..J) -> column after dims
        first = 3; last = 3 + len(RUBRIC_DIMENSIONS) - 1
        ov_col = last + 1
        from openpyxl.utils import get_column_letter as gcl
        sc.cell(row=i, column=ov_col, value=f"=AVERAGE({gcl(first)}{i}:{gcl(last)}{i})")
        for c in range(1, len(hdr) + 1):
            sc.cell(row=i, column=c).border = thin
    sc.column_dimensions["A"].width = 16; sc.column_dimensions["B"].width = 10
    for c in range(3, 3 + len(RUBRIC_DIMENSIONS)):
        from openpyxl.utils import get_column_letter as gcl
        sc.column_dimensions[gcl(c)].width = 11
    sc.column_dimensions[gcl(len(hdr))].width = 40
    sc.freeze_panes = "C2"
    wb.save(str(out_path))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--conditions", nargs="+", default=["no_refine", "datr_full"])
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--out", default=str(Path(__file__).resolve().parents[1] / "human_eval"))
    args = ap.parse_args()
    random.seed(args.seed)

    OUT = Path(args.out)
    d_lec = OUT / "01_selected_lectures"; d_lec.mkdir(parents=True, exist_ok=True)
    d_form = OUT / "02_rater_forms"; d_form.mkdir(parents=True, exist_ok=True)
    d_map = OUT / "04_rater_responses"; d_map.mkdir(parents=True, exist_ok=True)

    # 1) Selection
    selected = []   # (topic, model, condition, seed, ext, lecture)
    for topic in TOPICS:
        for model in MODELS:
            for cond in args.conditions:
                rows = load_cell(topic, model, cond)
                if not rows:
                    print(f"  [!] no data for {topic}/{model}/{cond}"); continue
                seed, ext, lec = pick_median(rows)
                selected.append((topic, model, cond, seed, ext, lec))

    # 2) Anonymization
    ids = [f"L{i+1:02d}" for i in range(len(selected))]
    order = list(range(len(selected))); random.shuffle(order)
    mapping = []
    print(f"\nSelected {len(selected)} lectures. Anonymizing and rendering...")
    for new_i, orig_i in enumerate(order):
        lid = ids[new_i]
        topic, model, cond, seed, ext, lec = selected[orig_i]
        mapping.append({"lecture_id": lid, "topic": topic, "model": model,
                        "condition": cond, "seed": seed, "panel_ext": round(ext, 3)})
        # 3) Render PDF
        try:
            render_pdf(lid, lec, d_lec / f"{lid}.pdf")
        except Exception as e:
            (d_lec / f"{lid}.md").write_text(lec, encoding="utf-8")
            print(f"   {lid}: PDF failed ({type(e).__name__}) — saved as .md")

    # Secret mapping file
    import csv
    mp = d_map / "_secret_mapping_RESEARCHER_ONLY.csv"
    with open(mp, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["lecture_id", "topic", "model", "condition", "seed", "panel_ext"])
        w.writeheader(); w.writerows(sorted(mapping, key=lambda x: x["lecture_id"]))

    # 4) Rating form
    build_rating_form(sorted(ids), d_form / "rater_form.xlsx")

    print(f"\n[done]:")
    print(f"   {len(ids)} anonymized lectures -> {d_lec}")
    print(f"   rating form              -> {d_form}/rater_form.xlsx")
    print(f"   secret mapping           -> {mp.name} (do NOT send to raters!)")


if __name__ == "__main__":
    main()
